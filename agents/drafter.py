"""
agents/drafter.py — Produce claim kit: editable .docx letter + document checklist.
All letters are clearly marked DRAFT. Nothing is ever sent or filed.
"""
from __future__ import annotations
import logging
import os
import uuid
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from models import (
    CaseProfile, TaskItem, ClaimKit, DocumentChecklist, ApprovalStatus
)
from utils.kb_loader import get_card

logger = logging.getLogger(__name__)

OUTPUT_DIR = "output"


def _ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def _draft_letter_text(
    profile: CaseProfile,
    task: TaskItem,
    today: str,
) -> str:
    """Generate the plain-text version of the draft letter."""
    institution = task.institution
    deceased = profile.deceased_name
    claimant = profile.claimant_name
    relationship = profile.claimant_relationship

    return f"""[DRAFT — REVIEW BEFORE USE — NOT YET SENT]

Date: {today}

To,
The Manager / Claims Department
{institution}
[Branch Address — Please fill in]

Subject: Claim / Notification regarding the account / policy of the late {deceased}

Dear Sir / Madam,

I, {claimant}, {relationship} of the late {deceased}, wish to inform you of
his / her passing on [Date of Death — please verify]. I am writing to initiate
the necessary claim / settlement / transmission process.

Details of the deceased:
  Full Name   : {deceased}
  Date of Death: [Please verify from death certificate]
  Account / Policy / Folio No.: [XXXXXXXX — please fill in the last 4 digits known]
  Address     : [Address — please fill in]

Details of claimant:
  Full Name   : {claimant}
  Relationship: {relationship}
  Contact     : [Phone — please fill in]
  Address     : [Address — please fill in]

I am enclosing the following documents for your reference:
{chr(10).join(f"  {i+1}. {doc}" for i, doc in enumerate(task.required_documents))}

I request you to kindly process this claim / transmission at the earliest and
advise me on any additional steps or documents required.

Please note: This is a draft prepared for review only. Verify all details with
the relevant institution before submission. Saath does not file, send or submit
anything on your behalf.

Yours sincerely,

{claimant}
[Signature]
[Date]
[Contact details]
"""


def _write_docx(letter_text: str, filename: str) -> str:
    """Write the letter to a .docx file and return the path."""
    _ensure_output_dir()
    path = os.path.join(OUTPUT_DIR, filename)
    doc = Document()

    # Draft watermark banner
    banner = doc.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = banner.add_run("⚠ DRAFT — REVIEW BEFORE USE — DO NOT SUBMIT AS-IS ⚠")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    doc.add_paragraph("")

    for line in letter_text.split("\n"):
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)

    doc.save(path)
    return path


def run_drafter(
    profile: CaseProfile,
    tasks: list[TaskItem],
) -> list[ClaimKit]:
    """Produce one ClaimKit per task that has a known institution."""
    kits: list[ClaimKit] = []
    today = date.today().strftime("%d %B %Y")

    # Only draft kits for tasks that are not pure legal referrals
    draftable_card_ids = {
        "bank_account_claim",
        "life_insurance_claim",
        "epf_pension_claim",
        "mutual_fund_transmission",
        "shares_demat_transmission",
    }

    for task in tasks:
        if task.source and task.source.card_id not in draftable_card_ids:
            continue
        if not task.institution:
            continue

        kit_id = f"KIT_{uuid.uuid4().hex[:6].upper()}"
        letter_text = _draft_letter_text(profile, task, today)
        filename = f"draft_{task.source.card_id if task.source else task.task_id}_{kit_id}.docx"

        try:
            _write_docx(letter_text, filename)
            logger.info("Drafter: wrote %s", filename)
        except Exception as e:
            logger.error("Drafter: failed to write docx for %s: %s", task.task_id, e)
            filename = f"DRAFT_{kit_id}.txt"

        card = get_card(task.source.card_id) if task.source else None
        checklist_items = card.get("required_documents", []) if card else task.required_documents
        optional_items = []

        kit = ClaimKit(
            kit_id=kit_id,
            institution=task.institution,
            asset_type=task.source.card_id if task.source else "general",
            letter_filename=filename,
            letter_text=letter_text,
            checklist=DocumentChecklist(
                institution=task.institution,
                items=checklist_items,
                optional_items=optional_items,
            ),
            is_draft=True,
            sent=False,
            approval_status=ApprovalStatus.PENDING,
        )
        kits.append(kit)

    logger.info("Drafter: produced %d claim kits", len(kits))
    return kits
